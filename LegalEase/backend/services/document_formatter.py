from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from fpdf import FPDF


BASE_DIR = Path(__file__).resolve().parents[2]
LOGO_PATH = BASE_DIR / "assets" / "logo.png"


def sanitize_text(text: str) -> str:
    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2022": "-",
        "\u00a0": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text.encode("latin-1", "replace").decode("latin-1")


def format_txt(text: str) -> bytes:
    return text.strip().encode("utf-8")


def _split_markdown(text: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line:
            blocks.append(("blank", ""))
        elif line.startswith("# "):
            blocks.append(("title", line[2:].strip()))
        elif re.match(r"^##\s+", line):
            blocks.append(("heading", re.sub(r"^##\s+", "", line)))
        elif re.match(r"^\d+\.\s+", line):
            blocks.append(("number", line))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:].strip()))
        elif line.startswith("> "):
            blocks.append(("note", line[2:].strip()))
        elif line.startswith("**") and line.endswith("**"):
            blocks.append(("bold", line.strip("*")))
        else:
            blocks.append(("paragraph", line.replace("**", "")))
    return blocks


def format_docx(text: str, doc_type: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    if LOGO_PATH.exists():
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.5))

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    for kind, content in _split_markdown(text):
        if kind == "blank":
            continue
        if kind == "title":
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(content)
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(18)
        elif kind == "heading":
            p = document.add_paragraph()
            run = p.add_run(content)
            run.bold = True
            run.font.size = Pt(13)
        elif kind == "bullet":
            document.add_paragraph(content, style="List Bullet")
        elif kind == "number":
            document.add_paragraph(content)
        elif kind == "note":
            p = document.add_paragraph()
            run = p.add_run(content)
            run.italic = True
        else:
            document.add_paragraph(content)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("LegalEase | AI-assisted draft | Review with a qualified legal professional")
    fr.font.size = Pt(8)

    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class LegalEasePDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = doc_type
        self.set_auto_page_break(auto=True, margin=18)

    def header(self):
        if LOGO_PATH.exists():
            try:
                self.image(str(LOGO_PATH), x=90, y=8, w=30)
                self.ln(18)
            except Exception:
                self.ln(5)
        else:
            self.ln(5)
        self.set_font("Helvetica", "B", 8)
        self.cell(0, 5, "LegalEase", align="L")
        self.ln(4)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "", 7)
        self.cell(0, 6, "LegalEase | AI-assisted draft | Review before legal use", align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = LegalEasePDF(doc_type)
    pdf.set_margins(18, 25, 18)
    pdf.add_page()

    for kind, content in _split_markdown(sanitize_text(text)):
        if kind == "blank":
            pdf.ln(3)
        elif kind == "title":
            pdf.set_font("Helvetica", "B", 16)
            pdf.multi_cell(0, 8, content, align="C")
            pdf.ln(3)
        elif kind == "heading":
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(0, 7, content)
            pdf.ln(1)
        elif kind == "bullet":
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 6, "- " + content)
        elif kind == "number":
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 6, content)
        elif kind == "note":
            pdf.set_font("Helvetica", "I", 8)
            pdf.multi_cell(0, 5, content)
        else:
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 6, content)

    output = pdf.output()
    return bytes(output)
